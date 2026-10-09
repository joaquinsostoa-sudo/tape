"""Censo Agropecuario 2022: fincas y superficie por distrito y estrato de tamaño (17 departamentos, sin Capital).

Fuente: data/raw/censo_agropecuario/CAN2022_fincas_por_distrito.xlsx (una hoja por departamento). Cada hoja trae el
total del departamento y, después, un bloque por distrito: encabezado con el código de 4 cifras (departamento + distrito),
su cantidad de fincas y su superficie, seguido de los estratos de tamaño de finca. Para anonimizar, los estratos
extremos se agregan en algunos distritos ("De 1.000 y más", "(*) Estrato agregado"): los estratos no son comparables
entre distritos, los totales sí. Los distritos sin fincas vienen con superficie vacía.
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[2]
XLSX = RAIZ / "data" / "raw" / "censo_agropecuario" / "CAN2022_fincas_por_distrito.xlsx"
VOLUMEN_I = RAIZ / "data" / "raw" / "censo_agropecuario" / "CAN2022_Volumen_I.xlsx"
CLEAN = RAIZ / "data" / "clean"
RE_DISTRITO = re.compile(r"^(\d{4})\s+(.+)$")
RE_DEPARTAMENTO = re.compile(r"^(\d{2})\.\s+(.+)$")
RE_TITULO = re.compile(r"^C(\d+):\s*([^.]+)\.")


def _num(v: object) -> float | None:
    return None if v in (None, "", "-") else float(v)


def _limpiar(s: str) -> str:
    return " ".join(str(s).split())


def leer(ruta: Path = XLSX) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(totales, estratos): totales por departamento y distrito; estratos en formato largo."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    totales, estratos = [], []
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        titulo = RE_TITULO.match(str(filas[1][0]))
        if titulo is None:
            raise ValueError(f"hoja {ws.title!r}: título no reconocido: {filas[1][0]!r}")
        dep, dep_nombre = int(titulo.group(1)), _limpiar(titulo.group(2))
        actual: dict | None = None
        for r in filas[4:]:
            etiqueta = r[0]
            if etiqueta is None or str(etiqueta).startswith(("(*)", "Fuente")):
                continue
            etiqueta = _limpiar(etiqueta)
            d = RE_DISTRITO.match(etiqueta)
            if d and d.group(1)[:2] == f"{dep:02d}":
                actual = {"dep_codigo": dep, "departamento": dep_nombre, "distrito_codigo": d.group(1),
                          "distrito": d.group(2)}
            elif RE_DEPARTAMENTO.match(etiqueta) or etiqueta == "Total":
                actual = {"dep_codigo": dep, "departamento": dep_nombre, "distrito_codigo": None, "distrito": None}
            elif actual is None:
                continue
            else:  # fila de estrato de tamaño
                estratos.append({**actual, "estrato": etiqueta, "fincas": _num(r[1]) or 0.0,
                                 "superficie_ha": _num(r[2]) or 0.0})
                continue
            totales.append({**actual, "fincas": _num(r[1]) or 0.0, "superficie_ha": _num(r[2]) or 0.0})
    return pd.DataFrame(totales), pd.DataFrame(estratos)


def totales_volumen_i(ruta: Path = VOLUMEN_I) -> pd.DataFrame:
    """Totales por departamento de 2022 del Cuadro 1 del Volumen I, para contrastar."""
    ws = openpyxl.load_workbook(ruta, read_only=True, data_only=True)["Cuadro 1"]
    salida = []
    for r in ws.iter_rows(values_only=True):
        m = RE_DEPARTAMENTO.match(_limpiar(r[0])) if r[0] else None
        if m:
            salida.append({"dep_codigo": int(m.group(1)), "fincas": float(r[3]), "superficie_ha": float(r[4])})
    return pd.DataFrame(salida)


# Nombre de la ciudad en la tabla del mapa MIC (capa A) para los distritos del CAN que se escriben distinto.
# Clave: código de distrito del censo (4 cifras). Los demás se unen por nombre normalizado y departamento.
ALIAS_MIC = {
    "0105": "SAN CARLOS", "0107": "YBY YAU", "0108": "AZOTEY", "0109": "JOSE FELIX LOPEZ",
    "0201": "SAN PEDRO", "0214": "GRAL. ISIDORO RESQUIN", "0216": "GUAYAIBI",
    "0403": "CPTAN. MAURICIO JOSE TROCHE", "0406": "GENERAL EUGENIO A. GARAY", "0407": "COLONIA INDEPENDENCIA",
    "0416": "DR. BOTTRELL", "0505": "STA. ROSA DEL MBUTUY", "0516": "MCAL. FRANCISCO SOLANO LOPEZ",
    "0604": "DR. MOISES BERTONI", "0716": "LEANDRO OVIEDO", "0718": "MAYOR OTANO",
    "0801": "SAN JUAN BAUTISTA", "0904": "GRAL.BERNARDINO CABALLERO", "0912": "SAN ROQUE G. DE SANTA CRUZ",
    "0914": "TEBICUARYMI", "0917": "YBYTIMI", "1210": "MAYOR JOSE D. MARTINEZ", "1305": "KARAPA'I",
    "1403": "CURUGUATY", "1406": "YPEHU", "1407": "GRAL. FRANCISCO CABALLERO ALVAREZ", "1409": "LA PALOMA",
    "1503": "PTO. PINASCO", "1507": "TTE. IRALA FERNANDEZ", "1508": "TTE. ESTEBAN MARTINEZ",
    "1509": "GRAL. BRUGUEZ", "1602": "MCAL. ESTIGARRIBIA",
}


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    return " ".join("".join(ch if ch.isalnum() else " " for ch in s).split())


def unir_ciudades(totales: pd.DataFrame, ciudades: pd.DataFrame) -> pd.DataFrame:
    """Las 263 ciudades del mapa MIC con las fincas y la superficie del CAN. Las que no figuran en el CAN
    (`en_can` falso) quedan con valores vacíos, no cero."""
    d = totales[totales.distrito_codigo.notna()].copy()
    d["clave"] = [_norm(ALIAS_MIC.get(cod, nom)) for cod, nom in zip(d.distrito_codigo, d.distrito, strict=True)]
    c = ciudades.assign(clave=ciudades.ciudad.map(_norm))
    m = c.merge(d[["dep_codigo", "clave", "distrito_codigo", "fincas", "superficie_ha"]],
                on=["dep_codigo", "clave"], how="left", validate="1:1", indicator=True)
    sobran = set(d.distrito_codigo) - set(m.distrito_codigo.dropna())
    if sobran:
        raise ValueError(f"distritos del CAN sin ciudad en el mapa MIC: {sorted(sobran)}")
    m["en_can"] = m._merge == "both"
    m["superficie_media_ha"] = m.superficie_ha / m.fincas.where(m.fincas > 0)
    m["fincas_por_1000_hab"] = 1000 * m.fincas / m.poblacion_2022
    return m.drop(columns=["clave", "_merge"])


def main() -> None:
    totales, estratos = leer()
    totales.to_csv(CLEAN / "can_fincas_totales.csv", index=False, encoding="utf-8")
    estratos.to_csv(CLEAN / "can_fincas_estratos.csv", index=False, encoding="utf-8")
    dist = totales[totales.distrito_codigo.notna()]
    print(f"{totales.dep_codigo.nunique()} departamentos; {len(dist)} distritos; {int(dist.fincas.sum()):,} fincas; "
          f"{dist.superficie_ha.sum():,.0f} ha")
    ciudades = CLEAN / "ciudades_mic.csv"
    if ciudades.exists():
        u = unir_ciudades(totales, pd.read_csv(ciudades))
        u.to_csv(CLEAN / "can_fincas_distrito.csv", index=False, encoding="utf-8")
        print(f"unión con el mapa MIC: {int(u.en_can.sum())} de {len(u)} ciudades; sin fila en el CAN: "
              f"{u.loc[~u.en_can, 'ciudad'].tolist()}")


if __name__ == "__main__":
    main()
