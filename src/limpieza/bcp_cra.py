"""Cuentas Regionales Anuales (CRA) del BCP: PIB por rama y departamento, 2021-2024, base 2014.

Lee las 18 hojas departamentales del anexo (00 Asunción ... 17 Alto Paraguay) y las deja en formato
largo: una fila por (departamento, año, rama, medida). Medidas: nominal, real y deflactor (índice) y
sus variaciones interanuales. Unidad: millones de guaraníes (nominal y real); el deflactor es un índice.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[2]
EXCEL = RAIZ / "data" / "raw" / "bcp_cuentas_regionales" / "Anexo_Estadistico_CRA.xlsx"
SALIDA = RAIZ / "data" / "clean" / "pib_regional.parquet"
MEDIDAS = {"Nominal": "nominal", "Real": "real", "Deflactor": "deflactor"}
RAMAS = {
    "Agricultura": "agricultura", "Ganadería, forestal, pesca y minería": "ganaderia_forestal_pesca_mineria",
    "Manufactura": "manufactura", "Electricidad y agua": "electricidad_agua", "Construcción": "construccion",
    "Servicios": "servicios", "Total Valor Agregado Sectorial": "valor_agregado_total",
    "Impuestos a los Productos": "impuestos_productos", "Producto Interno Bruto": "pib",
}


def _texto(c: object) -> str:
    return c.strip() if isinstance(c, str) else ""


def leer_hoja(ws: openpyxl.worksheet.worksheet.Worksheet) -> pd.DataFrame:
    """Una hoja departamental -> tabla larga."""
    filas = list(ws.iter_rows(values_only=True))
    i_med = next(i for i, r in enumerate(filas) if "Nominal" in [_texto(c) for c in r])
    i_ramas, i_bloque = i_med - 1, i_med - 2
    inicio_var = next((j for j, c in enumerate(filas[i_bloque]) if "Variaci" in _texto(c)), len(filas[i_med]))
    cod, nombre = re.match(r"\s*(\d\d)\.\s*(.+?)\s*$", _texto(filas[5][0])).groups()
    columnas: dict[int, tuple[str, str]] = {}
    rama = ""
    for j, c in enumerate(filas[i_med]):
        if _texto(filas[i_ramas][j]):
            rama = RAMAS[_texto(filas[i_ramas][j])]
        medida = MEDIDAS.get(_texto(c))
        if medida:
            columnas[j] = (rama, f"var_{medida}" if j >= inicio_var else medida)
    salida = []
    for r in filas[i_med + 1:]:
        etiqueta = str(r[0]).strip() if r[0] is not None else ""
        m = re.fullmatch(r"(\d{4})(\*?)", etiqueta)
        if not m:
            if salida:
                break
            continue
        for j, (rama, medida) in columnas.items():
            if isinstance(r[j], (int, float)):
                salida.append((cod, nombre, int(m.group(1)), bool(m.group(2)), rama, medida, float(r[j])))
    return pd.DataFrame(salida, columns=["depto_codigo", "departamento", "anio", "preliminar", "rama", "medida", "valor"])


def cargar(ruta: Path = EXCEL) -> pd.DataFrame:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    hojas = [wb[n] for n in wb.sheetnames if re.match(r"\d\d\.", n)]
    return pd.concat([leer_hoja(ws) for ws in hojas], ignore_index=True)


def main() -> None:
    df = cargar()
    df.to_parquet(SALIDA, index=False)
    print(f"{df.depto_codigo.nunique()} departamentos, años {df.anio.min()}-{df.anio.max()}, filas {len(df)}")


if __name__ == "__main__":
    main()
